import os
from pptx import Presentation
from pptx.util import Inches

def construire_presentation_avec_trame(graphiques, chemin_modele='export/templates/trame.pptx', chemin_sortie="export/presentations/Rapport_CODIR.pptx"):
    
    if not os.path.exists(chemin_modele):
        raise FileNotFoundError(f"Erreur : La trame '{chemin_modele}' est introuvable.")
        
    prs = Presentation(chemin_modele)
    print(f"Trame '{chemin_modele}' chargée avec succès.")

    for slide_data in graphiques:
        titre = slide_data.get('titre', 'Sans Titre')
        chemin_image = slide_data.get('image', None)
        index_masque = slide_data.get('index_masque', 0) 

        titre_mois = slide_data.get('titre_mois', None)
        
        try:
            layout = prs.slide_layouts[index_masque]
        except IndexError:
            print(f"Attention: Le masque {index_masque} n'existe pas. Utilisation du masque par défaut (0).")
            layout = prs.slide_layouts[0]

        slide = prs.slides.add_slide(layout)
        
        if slide.shapes.title:
            slide.shapes.title.text = titre

        # CORRECTION 1 : Recherche tolérante pour le titre du mois
        if titre_mois:
            for shape in slide.shapes:
                if "titre_mois" in shape.name:
                    # On cible directement le cadre de texte pour préserver la structure
                    if shape.has_text_frame:
                        shape.text_frame.text = titre_mois
                    break

        # CORRECTION 2 : Ajustement de l'image sans rognage
        if chemin_image and os.path.exists(chemin_image):
            image_inseree = False
            
            for shape in slide.shapes:
                is_img_chart = "IMG_CHART" in shape.name
                is_picture_placeholder = shape.is_placeholder and shape.placeholder_format.type == 18
                
                if is_img_chart or is_picture_placeholder:
                    # 1. On récupère les coordonnées de l'espace réservé
                    gauche = shape.left
                    haut = shape.top
                    largeur = shape.width
                    
                    # 2. On supprime l'espace réservé d'origine pour éviter le rognage forcé
                    try:
                        sp = shape.element
                        sp.getparent().remove(sp)
                    except Exception:
                        pass
                    
                    # 3. On ajoute l'image avec la largeur exacte (la hauteur s'adapte automatiquement)
                    slide.shapes.add_picture(chemin_image, left=gauche, top=haut, width=largeur)
                    image_inseree = True
                    break
            
            # Solution de repli
            if not image_inseree:
                print(f"Espace image introuvable sur la slide '{titre}'. Positionnement manuel.")
                slide.shapes.add_picture(chemin_image, Inches(1.0), Inches(1.5), width=Inches(11.3))
                
    # Création du dossier de sortie s'il n'existe pas
    dossier_sortie = os.path.dirname(chemin_sortie)
    if dossier_sortie and not os.path.exists(dossier_sortie):
        os.makedirs(dossier_sortie)

    # Sauvegarde du fichier final
    prs.save(chemin_sortie)
    print(f"✅ Présentation générée avec succès : {chemin_sortie}")